import os
import re
import logging
from typing import Dict, Any, Optional, Tuple
from docx import Document
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm

from config import IMAGE_SLOTS

logger = logging.getLogger("doc_engine")

class DocumentEngine:
    """Advanced OpenXML / Word Document Engine supporting Jinja2 templating,
    millimeter image scaling, and XML split-run token sanitization."""

    @staticmethod
    def process_document(
        template_path: str,
        output_path: str,
        context: Dict[str, Any],
        image_map: Optional[Dict[str, Dict[str, Any]]] = None,
        custom_find_replace: Optional[Dict[str, Any]] = None
    ) -> Tuple[bool, str]:
        """Processes a single DOCX template with context data, image substitutions, and custom word replacement.
        
        Args:
            template_path: Path to input .docx template.
            output_path: Destination path for rendered .docx.
            context: Dictionary of 78 canonical SMAP variables.
            image_map: Optional dict mapping tag_name -> { "path": str, "width_mm": int, "height_mm": int }
            custom_find_replace: Optional dict with {"find_text": str, "replace_text": str, "case_sensitive": bool}
            
        Returns:
            Tuple[bool, str]: (Success status, Error message or output path)
        """
        if not os.path.exists(template_path):
            return False, f"Template file not found: {template_path}"

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        try:
            # 1. Initialize DocxTemplate
            doc = DocxTemplate(template_path)
            render_context = dict(context)
            
            # Alias for Word templates that use {{hari_tgl_terbilang_ba_pemusnahan}}
            if "haritgl_terbilang_ba_pemusnahan" in render_context:
                render_context["hari_tgl_terbilang_ba_pemusnahan"] = render_context["haritgl_terbilang_ba_pemusnahan"]
                
            # Alias for Word templates that use {{tgl_berlaku_short}}
            if "tgl_berlaku_pendek" in render_context:
                render_context["tgl_berlaku_short"] = render_context["tgl_berlaku_pendek"]

            # Alias for Word templates that use {{jenis_usaha}} or {{jenis_usaha_bu}}
            if "jenis_usaha" in render_context:
                render_context["jenis_usaha_bu"] = render_context["jenis_usaha"]
            elif "jenis_usaha_bu" in render_context:
                render_context["jenis_usaha"] = render_context["jenis_usaha_bu"]

            
            # Create structured lists for Jinja2 dynamic table rows (e.g. {% tr for item in serti_standar_list %})
            serti_list = []
            sbu_list = []
            for i in range(1, 51):
                serti_val = render_context.get(f"serti_standar_{i}")
                if serti_val and str(serti_val).strip() and str(serti_val).strip() != "-":
                    serti_list.append({
                        "index": i,
                        "nomor": str(serti_val).strip(),
                        "tgl": str(render_context.get(f"tgl_serti_standar_{i}", "")).strip()
                    })
                
                sbu_val = render_context.get(f"sbu_umku_{i}")
                if sbu_val and str(sbu_val).strip() and str(sbu_val).strip() != "-":
                    sbu_list.append({
                        "index": i,
                        "nomor": str(sbu_val).strip(),
                        "tgl": str(render_context.get(f"tgl_sbu_umku_{i}", "")).strip()
                    })
            
            render_context["serti_standar_list"] = serti_list
            render_context["sbu_umku_list"] = sbu_list

            # 2. Convert uploaded images to InlineImage objects
            if image_map:
                for tag_name, img_info in image_map.items():
                    img_path = img_info.get("path")
                    if img_path and os.path.exists(img_path):
                        h_mm = img_info.get("height_mm") or IMAGE_SLOTS.get(tag_name, {}).get("default_height_mm", 25)
                        try:
                            inline_img = InlineImage(
                                doc,
                                img_path,
                                height=Mm(h_mm)
                            )
                            render_context[tag_name] = inline_img
                        except Exception as img_err:
                            logger.warning(f"Failed to load InlineImage for '{tag_name}': {img_err}")

            # 3. Render Jinja2 template tags across Body, Headers, Footers, and Tables
            doc.render(render_context)
            doc.save(output_path)

            # 4. Fallback Pass: Raw XML split-run sanitizer for legacy non-jinja {token} or ${token}
            DocumentEngine.sanitize_legacy_tokens(output_path, render_context)

            # 5. Alt Text Image Replacement Pass
            if image_map:
                DocumentEngine.replace_images_by_alt_text(output_path, image_map)

            # 6. Custom Word/Phrase Find & Replace Pass (supports multiple Form 1 -> Form 2 rules)
            if custom_find_replace:
                pairs_to_run = []
                if isinstance(custom_find_replace, list):
                    pairs_to_run = custom_find_replace
                elif isinstance(custom_find_replace, dict):
                    if "pairs" in custom_find_replace and isinstance(custom_find_replace["pairs"], list):
                        default_cs = custom_find_replace.get("case_sensitive", True)
                        for item in custom_find_replace["pairs"]:
                            p_copy = dict(item)
                            if "case_sensitive" not in p_copy:
                                p_copy["case_sensitive"] = default_cs
                            pairs_to_run.append(p_copy)
                    elif custom_find_replace.get("find_text"):
                        pairs_to_run = [custom_find_replace]

                for p_item in pairs_to_run:
                    f_text = p_item.get("find_text", "")
                    if f_text:
                        r_text = p_item.get("replace_text", "")
                        cs = p_item.get("case_sensitive", True)
                        DocumentEngine.replace_text_in_document(
                            docx_path=output_path,
                            find_text=f_text,
                            replace_text=r_text,
                            case_sensitive=cs
                        )

            return True, output_path

        except Exception as e:
            logger.error(f"Error processing document '{template_path}': {e}", exc_info=True)
            return False, str(e)

    @staticmethod
    def sanitize_legacy_tokens(docx_path: str, context: Dict[str, Any]) -> None:
        """Traverses paragraphs, table cells, headers, and footers in a docx file
        and performs raw text replacement for legacy {variable} and ${variable} tags,
        fixing Word XML split-run issues.
        """
        try:
            doc = Document(docx_path)
            
            # Sanitizer helper for a list of paragraphs
            def sanitize_paragraphs(paragraphs):
                for p in paragraphs:
                    full_text = p.text
                    if not full_text:
                        continue
                    
                    modified = False
                    for key, val in context.items():
                        if str(val) is None:
                            continue
                        str_val = str(val)
                        
                        patterns = [
                            f"{{{key}}}",
                            f"${{{key}}}",
                            f"{{{{{key}}}}}"  # Jinja tag unrendered fallback
                        ]
                        for pat in patterns:
                            if pat in full_text:
                                full_text = full_text.replace(pat, str_val)
                                modified = True
                    
                    if modified and p.runs:
                        p.runs[0].text = full_text
                        for r in p.runs[1:]:
                            r.text = ""

            # 1. Process main body paragraphs
            sanitize_paragraphs(doc.paragraphs)

            # 2. Process all tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        sanitize_paragraphs(cell.paragraphs)

            # 3. Process Headers & Footers across all sections
            for section in doc.sections:
                if section.header:
                    sanitize_paragraphs(section.header.paragraphs)
                    for h_table in section.header.tables:
                        for h_row in h_table.rows:
                            for h_cell in h_row.cells:
                                sanitize_paragraphs(h_cell.paragraphs)
                if section.footer:
                    sanitize_paragraphs(section.footer.paragraphs)
                    for f_table in section.footer.tables:
                        for f_row in f_table.rows:
                            for f_cell in f_row.cells:
                                sanitize_paragraphs(f_cell.paragraphs)

            doc.save(docx_path)
        except Exception as e:
            logger.warning(f"Legacy XML split-run sanitizer encountered minor issue: {e}")

    @staticmethod
    def replace_images_by_alt_text(docx_path: str, image_map: Dict[str, Dict[str, Any]]) -> None:
        """Replaces images in the document based on their Alt Text matching tags in the image_map.
        Preserves image placement, sizing, and styling perfectly."""
        try:
            doc = Document(docx_path)
            replaced_count = 0
            
            # Alt text to image path mapping
            alt_to_path = {}
            for tag, info in image_map.items():
                if info.get("path") and os.path.exists(info["path"]):
                    alt_to_path[tag.lower().strip()] = info["path"]
                    
            if not alt_to_path:
                return

            valid_content_types = (
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml',
                'application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml',
                'application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml'
            )

            for part in doc.part.package.parts:
                if part.content_type not in valid_content_types:
                    continue
                    
                for element in part.element.iter():
                    # Look for <wp:docPr> which contains the alt text in 'descr', 'title', or 'name' attribute
                    if element.tag.endswith('}docPr'):
                        descr = (element.get('descr') or "").lower().strip()
                        title = (element.get('title') or "").lower().strip()
                        name = (element.get('name') or "").lower().strip()
                        
                        matched_tag = None
                        if descr in alt_to_path:
                            matched_tag = descr
                        elif title in alt_to_path:
                            matched_tag = title
                        elif name in alt_to_path:
                            matched_tag = name
    
                        if matched_tag:
                            new_image_path = alt_to_path[matched_tag]
                            
                            # Find the parent drawing or pict element
                            parent = element.getparent()
                            while parent is not None and not parent.tag.endswith('}drawing') and not parent.tag.endswith('}pict'):
                                parent = parent.getparent()
                                
                            if parent is not None:
                                # Look for a:blip
                                for blip in parent.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
                                    rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                                    if rId:
                                        # Replace the binary blob of the image part
                                        try:
                                            rel = part.rels[rId]
                                            
                                            # Use PIL to read the user's image and convert it to match the DOCX part's content type
                                            from PIL import Image
                                            import io
                                            img = Image.open(new_image_path)
                                            
                                            target_format = "PNG"
                                            c_type = rel.target_part.content_type.lower()
                                            if "jpeg" in c_type or "jpg" in c_type:
                                                target_format = "JPEG"
                                                if img.mode in ("RGBA", "P"):
                                                    img = img.convert("RGB")
                                            elif "gif" in c_type:
                                                target_format = "GIF"
                                            elif "bmp" in c_type:
                                                target_format = "BMP"
                                                
                                            out_io = io.BytesIO()
                                            img.save(out_io, format=target_format)
                                            
                                            rel.target_part._blob = out_io.getvalue()
                                            replaced_count += 1
                                            logger.info(f"Replaced image with alt text '{matched_tag}' (converted to {target_format})")
                                        except Exception as rel_err:
                                            logger.warning(f"Failed to replace image part for '{matched_tag}': {rel_err}")

            if replaced_count > 0:
                doc.save(docx_path)
        except Exception as e:
            logger.error(f"Error in replace_images_by_alt_text: {e}")

    @staticmethod
    def highlight_audit_month(docx_path: str, target_month_idx: int) -> None:
        """
        Scans tables in the document for the Audit Program table (identified by month headers)
        and highlights the column corresponding to the target_month_idx (1-12) in yellow,
        while removing the background color from all other month columns.
        """
        try:
            from docx.oxml.ns import nsdecls
            from docx.oxml import parse_xml
            from docx import Document
            
            doc = Document(docx_path)
            
            # Map any known header string to its month index 1-12
            month_header_map = {
                "jan": 1, "januari": 1,
                "feb": 2, "februari": 2,
                "mar": 3, "maret": 3,
                "apr": 4, "april": 4,
                "may": 5, "mei": 5,
                "jun": 6, "juni": 6,
                "jul": 7, "juli": 7,
                "aug": 8, "agst": 8, "ags": 8, "agustus": 8,
                "sep": 9, "sept": 9, "september": 9,
                "oct": 10, "okt": 10, "oktober": 10,
                "nov": 11, "november": 11,
                "dec": 12, "des": 12, "desember": 12
            }
            
            modified = False
            for table in doc.tables:
                is_audit_table = False
                header_row_idx = 0
                month_col_map = {} 
                
                # Check the first two rows to find the headers
                for i, row in enumerate(table.rows[:2]):
                    for j, cell in enumerate(row.cells):
                        text = cell.text.lower().replace("\n", "").strip()
                        if text in month_header_map:
                            is_audit_table = True
                            month_col_map[j] = month_header_map[text]
                    if is_audit_table:
                        header_row_idx = i
                        break
                        
                if is_audit_table:
                    # Target color (Gold/Yellow)
                    highlight_color = "FFC000"
                    
                    # Apply coloring to all data rows
                    for row in table.rows[header_row_idx+1:]:
                        for col_idx, m_idx in month_col_map.items():
                            if col_idx < len(row.cells):
                                cell = row.cells[col_idx]
                                tcPr = cell._tc.get_or_add_tcPr()
                                
                                # Remove existing shading
                                tcPr_shd = tcPr.find(r'{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
                                if tcPr_shd is not None:
                                    tcPr.remove(tcPr_shd)
                                    modified = True
                                    
                                # Apply new shading if it's the target month
                                if m_idx == target_month_idx:
                                    shd = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), highlight_color))
                                    tcPr.append(shd)
                                    modified = True
                                    
            if modified:
                doc.save(docx_path)
        except Exception as e:
            logger.error(f"Error in highlight_audit_month: {e}", exc_info=True)

    @staticmethod
    def replace_text_in_document(
        docx_path: str,
        find_text: str,
        replace_text: str,
        case_sensitive: bool = True
    ) -> int:
        """Searches and replaces all occurrences of find_text with replace_text across body
        paragraphs, tables, headers, and footers in a Word (.docx) file.
        Correctly handles split-run Word XML tokens and supports Case Sensitive / Non-Case Sensitive modes.
        """
        if not find_text or not os.path.exists(docx_path):
            return 0

        try:
            doc = Document(docx_path)

            def replace_in_str(text: str) -> str:
                if case_sensitive:
                    return text.replace(find_text, replace_text)
                else:
                    return re.sub(re.escape(find_text), lambda _: replace_text, text, flags=re.IGNORECASE)

            def is_match(text: str) -> bool:
                if not find_text or not text:
                    return False
                if case_sensitive:
                    return find_text in text
                else:
                    return bool(re.search(re.escape(find_text), text, re.IGNORECASE))

            def process_paragraphs(paragraphs) -> int:
                count = 0
                for p in paragraphs:
                    if not p.text or not is_match(p.text):
                        continue

                    # 1. Try replacing within individual runs first to preserve local formatting
                    single_run_matched = False
                    for r in p.runs:
                        if is_match(r.text):
                            r.text = replace_in_str(r.text)
                            single_run_matched = True
                            count += 1

                    # 2. If token is split across multiple runs in XML, merge into first run
                    if not single_run_matched and is_match(p.text):
                        full_text = replace_in_str(p.text)
                        if p.runs:
                            p.runs[0].text = full_text
                            for r in p.runs[1:]:
                                r.text = ""
                        else:
                            p.text = full_text
                        count += 1
                return count

            def process_table(table) -> int:
                cnt = 0
                for row in table.rows:
                    for cell in row.cells:
                        cnt += process_paragraphs(cell.paragraphs)
                        for nested_t in cell.tables:
                            cnt += process_table(nested_t)
                return cnt

            total_replaced = process_paragraphs(doc.paragraphs)
            for t in doc.tables:
                total_replaced += process_table(t)

            for section in doc.sections:
                for header in [section.header, getattr(section, 'first_page_header', None), getattr(section, 'even_page_header', None)]:
                    if header:
                        total_replaced += process_paragraphs(header.paragraphs)
                        for t in header.tables:
                            total_replaced += process_table(t)
                for footer in [section.footer, getattr(section, 'first_page_footer', None), getattr(section, 'even_page_footer', None)]:
                    if footer:
                        total_replaced += process_paragraphs(footer.paragraphs)
                        for t in footer.tables:
                            total_replaced += process_table(t)

            if total_replaced > 0:
                doc.save(docx_path)
                logger.info(f"Replaced {total_replaced} occurrence(s) of '{find_text}' in '{os.path.basename(docx_path)}' (Case Sensitive: {case_sensitive})")

            return total_replaced
        except Exception as e:
            logger.error(f"Error in replace_text_in_document for '{docx_path}': {e}", exc_info=True)
            return 0

