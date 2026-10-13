import os
import sys
import shutil
import tempfile
import unittest
from docx import Document

from config import DEFAULT_VARIABLE_VALUES, ALL_SMAP_VARIABLES
from core.doc_engine import DocumentEngine
from core.pdf_engine import PDFEngine
from core.profile_manager import ProfileManager
from core.batch_worker import BatchProcessor

class TestSMAPAutoDocEngine(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="smap_test_")
        self.assets_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
        self.sample_logo = os.path.join(self.assets_dir, "sample_logo.png")
        self.sample_stamp = os.path.join(self.assets_dir, "sample_stamp.png")

    def tearDown(self):
        if os.path.exists(self.test_dir):
            shutil.rmtree(self.test_dir, ignore_errors=True)

    def create_sample_docx_template(self, docx_path: str):
        doc = Document()
        doc.add_heading("DOKUMEN SPESIFIKASI SMAP ISO 37001", level=1)
        doc.add_paragraph("Nama BU: {{bu_caps}} ({{inisial_bu}})")
        doc.add_paragraph("Alamat BU: {{alamat_bu}}, {{kab_kota}}")
        doc.add_paragraph("Direktur Utama: {{direktur}}")
        doc.add_paragraph("NPWP BU: {npwp_bu}")  # Raw legacy token
        doc.add_paragraph("No. KUMHAM: ${no_kumham}")  # Raw legacy token

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Variabel"
        table.cell(0, 1).text = "Nilai"
        table.cell(1, 0).text = "Ketua FKAP"
        table.cell(1, 1).text = "{{fkap}}"

        doc.save(docx_path)

    def test_doc_engine_and_xml_sanitizer(self):
        tpl_path = os.path.join(self.test_dir, "template.docx")
        out_path = os.path.join(self.test_dir, "output_rendered.docx")
        self.create_sample_docx_template(tpl_path)

        context = dict(DEFAULT_VARIABLE_VALUES)
        context["inisial_bu"] = "PT ABS"
        context["bu_caps"] = "PT ANUGERAH BERKAH SEJAHTERA"
        context["direktur"] = "Bambang Sudirjo"
        context["npwp_bu"] = "01.234.567.8-012.000"
        context["no_kumham"] = "AHU-0012345.AH.01.01"

        img_map = {}
        if os.path.exists(self.sample_logo):
            img_map["logo_bu"] = {"path": self.sample_logo, "height_mm": 20}

        success, res_path = DocumentEngine.process_document(tpl_path, out_path, context, img_map)
        self.assertTrue(success, f"Document processing failed: {res_path}")
        self.assertTrue(os.path.exists(out_path))

        # Verify text replacement
        res_doc = Document(out_path)
        full_text = "\n".join([p.text for p in res_doc.paragraphs])
        self.assertIn("PT ANUGERAH BERKAH SEJAHTERA", full_text)
        self.assertIn("Bambang Sudirjo", full_text)
        self.assertIn("01.234.567.8-012.000", full_text)
        self.assertIn("AHU-0012345.AH.01.01", full_text)

    def test_profile_manager_json_excel(self):
        json_path = os.path.join(self.test_dir, "profile.json")
        sample_profile = ProfileManager.get_sample_smap_profile()
        sample_profile["variables"]["inisial_bu"] = "PT ABS"
        sample_profile["variables"]["jenis_usaha"] = "Jasa Konsultansi Konstruksi"

        # Save & Load JSON
        ok_save, _ = ProfileManager.save_profile_json(sample_profile, json_path)
        self.assertTrue(ok_save)

        ok_load, loaded_profile, _ = ProfileManager.load_profile_json(json_path)
        self.assertTrue(ok_load)
        self.assertEqual(loaded_profile["variables"]["inisial_bu"], "PT ABS")
        self.assertEqual(loaded_profile["variables"]["jenis_usaha"], "Jasa Konsultansi Konstruksi")

        # Export & Import Excel
        excel_path = os.path.join(self.test_dir, "records.xlsx")
        ok_exp, _ = ProfileManager.export_profile_excel([sample_profile["variables"]], excel_path)
        self.assertTrue(ok_exp)

        ok_imp, records, _ = ProfileManager.import_profile_excel(excel_path)
        self.assertTrue(ok_imp)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["inisial_bu"], "PT ABS")
        self.assertEqual(records[0]["jenis_usaha"], "Jasa Konsultansi Konstruksi")

    def test_pdf_engine_watermark_merge(self):
        # Create small test PDF using PyMuPDF
        pdf1_path = os.path.join(self.test_dir, "doc1.pdf")
        pdf2_path = os.path.join(self.test_dir, "doc2.pdf")

        import fitz
        doc1 = fitz.open()
        page1 = doc1.new_page()
        page1.insert_text((50, 50), "Document 1 Content")
        doc1.save(pdf1_path)
        doc1.close()

        doc2 = fitz.open()
        page2 = doc2.new_page()
        page2.insert_text((50, 50), "Document 2 Content")
        doc2.save(pdf2_path)
        doc2.close()

        # Watermark Test
        wm_out = os.path.join(self.test_dir, "doc1_wm.pdf")
        ok_wm, _ = PDFEngine.apply_watermark(pdf1_path, wm_out, watermark_text="CONFIDENTIAL")
        self.assertTrue(ok_wm)
        self.assertTrue(os.path.exists(wm_out))

        # Stamp Test
        if os.path.exists(self.sample_stamp):
            st_out = os.path.join(self.test_dir, "doc1_st.pdf")
            ok_st, _ = PDFEngine.apply_stamp_or_signature(pdf1_path, st_out, self.sample_stamp)
            self.assertTrue(ok_st)
            self.assertTrue(os.path.exists(st_out))

        # Merge Test
        merged_out = os.path.join(self.test_dir, "merged.pdf")
        ok_merge, _ = PDFEngine.merge_pdfs([pdf1_path, pdf2_path], merged_out)
        self.assertTrue(ok_merge)
        self.assertTrue(os.path.exists(merged_out))

    def test_batch_worker(self):
        tpl_path = os.path.join(self.test_dir, "batch_template.docx")
        self.create_sample_docx_template(tpl_path)

        out_dir = os.path.join(self.test_dir, "batch_output")
        data_rows = [
            dict(DEFAULT_VARIABLE_VALUES, bu_folder="BU_Alpha", inisial_bu="Alpha"),
            dict(DEFAULT_VARIABLE_VALUES, bu_folder="BU_Beta", inisial_bu="Beta")
        ]

        processor = BatchProcessor(max_workers=2)
        summary = processor.process_batch(
            template_paths=[tpl_path],
            data_records=data_rows,
            output_dir=out_dir,
            convert_to_pdf=False
        )

        self.assertEqual(summary["total_tasks"], 2)
        self.assertEqual(summary["successful"], 2)

    def test_jenis_usaha_and_custom_find_replace(self):
        # 1. Test template with {{jenis_usaha}} placeholder
        doc = Document()
        doc.add_paragraph("Jenis Usaha: {{jenis_usaha}}")
        doc.add_paragraph("KATA KUNCI: PT Mitra Sejahtera.")
        doc.add_paragraph("KATA KUNCI KECIL: pt mitra sejahtera.")
        tpl_path = os.path.join(self.test_dir, "ju_template.docx")
        doc.save(tpl_path)

        out_path = os.path.join(self.test_dir, "ju_output.docx")
        context = dict(DEFAULT_VARIABLE_VALUES)
        context["jenis_usaha"] = "Jasa Konsultansi Konstruksi"

        # Case Sensitive Find & Replace
        cfr_cs = {
            "find_text": "PT Mitra Sejahtera",
            "replace_text": "PT Perkasa Abadi",
            "case_sensitive": True
        }
        ok, _ = DocumentEngine.process_document(tpl_path, out_path, context, custom_find_replace=cfr_cs)
        self.assertTrue(ok)

        res_doc = Document(out_path)
        text_full = "\n".join([p.text for p in res_doc.paragraphs])
        self.assertIn("Jenis Usaha: Jasa Konsultansi Konstruksi", text_full)
        self.assertIn("PT Perkasa Abadi", text_full)
        self.assertIn("pt mitra sejahtera", text_full)  # Lowercase was NOT replaced in CS mode!

        # Non-Case Sensitive Find & Replace
        out_ci_path = os.path.join(self.test_dir, "ju_output_ci.docx")
        cfr_ci = {
            "find_text": "PT Mitra Sejahtera",
            "replace_text": "PT Perkasa Abadi",
            "case_sensitive": False
        }
        ok_ci, _ = DocumentEngine.process_document(tpl_path, out_ci_path, context, custom_find_replace=cfr_ci)
        self.assertTrue(ok_ci)

        res_ci_doc = Document(out_ci_path)
        text_ci_full = "\n".join([p.text for p in res_ci_doc.paragraphs])
        self.assertIn("Jenis Usaha: Jasa Konsultansi Konstruksi", text_ci_full)
        self.assertIn("PT Perkasa Abadi.", text_ci_full)
        self.assertNotIn("pt mitra sejahtera", text_ci_full)  # Lowercase WAS replaced in CI mode!

        # 2. Test BatchProcessor with custom_find_replace
        batch_out_dir = os.path.join(self.test_dir, "batch_cfr_out")
        processor = BatchProcessor(max_workers=2)
        summary = processor.process_batch(
            template_paths=[tpl_path],
            data_records=[context],
            output_dir=batch_out_dir,
            convert_to_pdf=False,
            custom_find_replace=cfr_cs
        )
        self.assertEqual(summary["successful"], 1)

    def test_bu_find_replace_ui_integration(self):
        from unittest.mock import MagicMock
        from app import SMAPAutoDocApp
        page = MagicMock()
        app = SMAPAutoDocApp(page)

        # 1. Verify find & replace controls exist in Data Badan Usaha
        self.assertTrue(hasattr(app, "txt_batch_find_word"))
        self.assertTrue(hasattr(app, "txt_batch_replace_word"))
        self.assertTrue(hasattr(app, "chk_batch_case_sensitive"))
        self.assertTrue(hasattr(app, "find_replace_fields"))

        # 2. Verify default row count is exactly 1
        self.assertEqual(len(app.find_replace_fields), 1)
        self.assertTrue(app.chk_batch_case_sensitive.value)

        # 3. Test dynamically adding rows (like Tambah Sertifikat Standar & SBU/UMKU)
        app.add_find_replace_row(update=False)
        self.assertEqual(len(app.find_replace_fields), 2)
        app.add_find_replace_row(update=False)
        self.assertEqual(len(app.find_replace_fields), 3)

        # Populate rows
        app.find_replace_fields[0]["find"].value = "Kata A"
        app.find_replace_fields[0]["replace"].value = "Ganti A"
        app.find_replace_fields[1]["find"].value = "Kata B"
        app.find_replace_fields[1]["replace"].value = "Ganti B"
        app.find_replace_fields[2]["find"].value = ""  # empty should be skipped
        app.find_replace_fields[2]["replace"].value = "Ganti C"

        pairs = app.get_all_find_replace_pairs()
        self.assertEqual(len(pairs), 2)
        self.assertEqual(pairs[0]["find_text"], "Kata A")
        self.assertEqual(pairs[0]["replace_text"], "Ganti A")
        self.assertEqual(pairs[1]["find_text"], "Kata B")
        self.assertEqual(pairs[1]["replace_text"], "Ganti B")

        # Test deleting row
        del_item = app.find_replace_fields[1]
        app.remove_find_replace_row(del_item)
        self.assertEqual(len(app.find_replace_fields), 2)
        self.assertNotIn(del_item, app.find_replace_fields)

        # 4. Verify tab switching
        app.switch_to_input_data_bu()
        self.assertEqual(app.nav_rail.selected_index, 0)
        self.assertEqual(app.smap_tabs.selected_index, 0)
        self.assertEqual(app.form_tabs.selected_index, 0)

        app.nav_to_process_documents()
        self.assertEqual(app.smap_tabs.selected_index, 1)

        # 5. Verify reset form recreates controls cleanly with default 1 row
        app.reset_form_data()
        self.assertTrue(hasattr(app, "txt_batch_find_word"))
        self.assertEqual(len(app.find_replace_fields), 1)

    def test_multi_pair_find_and_replace_engine(self):
        doc = Document()
        doc.add_paragraph("Teks awal: PT ABC bergerak di bidang Jasa Konstruksi dengan pimpinan Bapak Budi.")
        tpl_path = os.path.join(self.test_dir, "multi_fr_tpl.docx")
        out_path = os.path.join(self.test_dir, "multi_fr_out.docx")
        doc.save(tpl_path)

        context = dict(DEFAULT_VARIABLE_VALUES)
        custom_find_replace = {
            "pairs": [
                {"find_text": "PT ABC", "replace_text": "PT XYZ PERKASA", "case_sensitive": True},
                {"find_text": "Jasa Konstruksi", "replace_text": "Jasa Konsultansi Konstruksi", "case_sensitive": True},
                {"find_text": "Bapak Budi", "replace_text": "Ibu Siti", "case_sensitive": True}
            ]
        }

        ok, res = DocumentEngine.process_document(tpl_path, out_path, context, custom_find_replace=custom_find_replace)
        self.assertTrue(ok)

        res_doc = Document(out_path)
        full_text = "\n".join([p.text for p in res_doc.paragraphs])
        self.assertIn("PT XYZ PERKASA", full_text)
        self.assertIn("Jasa Konsultansi Konstruksi", full_text)
        self.assertIn("Ibu Siti", full_text)
        self.assertNotIn("PT ABC", full_text)
        self.assertNotIn("Jasa Konstruksi dengan", full_text)
        self.assertNotIn("Bapak Budi", full_text)

if __name__ == "__main__":
    unittest.main()
